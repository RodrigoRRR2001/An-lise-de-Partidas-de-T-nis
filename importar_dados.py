import pandas as pd
import sqlite3
import os
from datetime import datetime
from openpyxl import load_workbook
def aplicar_regras_logicas(df):
    """Preenche automaticamente campos baseado em regras do tênis"""
    df = df.copy()

    # Converter colunas booleanas para float (0.0/1.0)
    bool_cols = ['ace', 'primeiro_servico', 'falha_servico', 'devolucao_dentro', 
                 'break_point', 'subiu_rede']
    for col in bool_cols:
        if col in df.columns:
            df[col] = df[col].astype(float)
    
    # Converter a coluna 'placar' para string ANTES de começar a preencher
    if 'placar' in df.columns:
        df['placar'] = df['placar'].astype(str)
    
    # Regra 1: Se ace = 1, ponto_num = sacador, devolucao_dentro = 0, falha_servico = 0
    cond_ace = (df['ace'] == 1.0)
    df.loc[cond_ace, 'tipo_ponto'] = df.loc[cond_ace, 'tipo_ponto'].fillna("Ace")
    df.loc[cond_ace, 'golpe_vencedor'] = df.loc[cond_ace, 'golpe_vencedor'].fillna("Saque")
    df.loc[cond_ace, ['direcao_golpe', 'num_trocas', 'devolucao_dentro', 'falha_servico']] = [None, 1, 0, 0]
    df.loc[cond_ace, 'ponto_num'] = df.loc[cond_ace, 'servidor']

    # Regra 2: Se primeiro_servico = 1, falha_servico = 0
    cond_primeiro_servico = (df['primeiro_servico'] == 1.0)
    df.loc[cond_primeiro_servico, 'falha_servico'] = 0

    # Regra 3: Se devolucao_dentro = 0, ponto_num = servidor
    cond_devolucao_fora = (df['devolucao_dentro'] == 0.0)
    df.loc[cond_devolucao_fora, 'ponto_num'] = df.loc[cond_devolucao_fora, 'servidor']

    # Regra 4: Se falha_servico = 1, ponto_num = quem não estiver sacando (1 - servidor)
    cond_falha_servico = (df['falha_servico'] == 1.0)
    df.loc[cond_falha_servico, 'ponto_num'] = 1 - df.loc[cond_falha_servico, 'servidor']

    # Regra 5: Se num_trocas > 2, devolucao_dentro = 1
    cond_trocas = (df['num_trocas'] > 2) & (df['num_trocas'].notna())
    df.loc[cond_trocas, 'devolucao_dentro'] = 1

    # Regra 6: Lógica do placar
    placar_sacador = 0
    placar_receptor = 0
    placar_atual = "00-00"
    ultimo_sacador = None

    for idx, row in df.iterrows():
        # Reinicia o placar se o sacador mudou
        if ultimo_sacador is not None and row['servidor'] != ultimo_sacador:
            placar_sacador = 0
            placar_receptor = 0
            placar_atual = "00-00"
        
        ultimo_sacador = row['servidor']
        
        # Determina vencedor do ponto se não estiver definido
        if pd.isna(df.at[idx, 'ponto_num']):
            if row['devolucao_dentro'] == 0:
                df.at[idx, 'ponto_num'] = row['servidor']
            elif row['falha_servico'] == 1:
                df.at[idx, 'ponto_num'] = 1 - row['servidor']
        
        # Atualiza placar apenas se ponto_num estiver definido
        if not pd.isna(df.at[idx, 'ponto_num']):
            vencedor = df.at[idx, 'ponto_num']
            sacador = row['servidor']
            
            if vencedor == sacador:
                # Ponto para o sacador
                if placar_sacador == 0: placar_sacador = 15
                elif placar_sacador == 15: placar_sacador = 30
                elif placar_sacador == 30: placar_sacador = 40
                elif placar_sacador == 40:
                    if placar_receptor == 40: placar_sacador = 'ADV'
                    elif placar_receptor == 'ADV': placar_sacador, placar_receptor = 40, 40
                    else: placar_sacador = 'V'
            else:
                # Ponto para o receptor
                if placar_receptor == 0: placar_receptor = 15
                elif placar_receptor == 15: placar_receptor = 30
                elif placar_receptor == 30: placar_receptor = 40
                elif placar_receptor == 40:
                    if placar_sacador == 40: placar_receptor = 'ADV'
                    elif placar_sacador == 'ADV': placar_sacador, placar_receptor = 40, 40
                    else: placar_receptor = 'V'
            
            # Formata o placar
            if 'V' in [placar_sacador, placar_receptor]:
                placar_atual = "Game"
                placar_sacador = placar_receptor = 0
            else:
                placar_atual = f"{placar_sacador}-{placar_receptor}"
            
            df.at[idx, 'placar'] = placar_atual

    return df
def formatar_valor(valor):
    """Formata valores para o SQL, especialmente datas e valores nulos"""
    if pd.isna(valor):
        return None
    elif isinstance(valor, pd.Timestamp):
        return valor.strftime('%Y-%m-%d')
    return valor

def inserir_dados_manual(conn, df, tabela, colunas_tabela):
    """Método alternativo para inserção de dados que evita o problema do to_sql"""
    cursor = conn.cursor()
    
    # Verificar colunas existentes na tabela
    cursor.execute(f"PRAGMA table_info({tabela})")
    colunas_existentes = [col[1] for col in cursor.fetchall()]
    
    # Filtrar colunas do DataFrame para incluir apenas as que existem na tabela
    colunas_df = [col for col in df.columns if col in colunas_existentes]
    
    if not colunas_df:
        print(f"Nenhuma coluna do DataFrame existe na tabela {tabela}")
        return 0
    
    placeholders = ','.join(['?'] * len(colunas_df))
    query = f"INSERT INTO {tabela} ({','.join(colunas_df)}) VALUES ({placeholders})"
    
    inseridos = 0
    for _, row in df.iterrows():
        try:
            # Preparar valores formatados
            valores = [formatar_valor(row[col]) for col in colunas_df]
            cursor.execute(query, valores)
            inseridos += 1
        except Exception as e:
            print(f"Erro ao inserir linha: {row.to_dict()}")
            print(f"Erro: {str(e)}")
            conn.rollback()
    
    conn.commit()
    return inseridos

def enviar_dados():
    # Configurações
   excel_path = "dados_tenis.xlsx"
   db_path = "tenis_analises_db.db"
    
    try:
        # ===== 1. VALIDAÇÃO INICIAL =====
        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"Arquivo Excel não encontrado: {excel_path}")
        
        # ===== 2. CONEXÃO COM O BANCO =====
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        
        # ===== 3. CARREGAR DADOS DO EXCEL =====
        df_partidas = pd.read_excel(excel_path, sheet_name="Partidas")
        df_rallys = pd.read_excel(excel_path, sheet_name="Digitação")
        
        # Padroniza nomes de colunas
        df_rallys = df_rallys.rename(columns={
            "set": "set_num",
            "game": "game_num",
            "ponto": "ponto_num"
        })
        
        # ===== 4. APLICA REGRAS LÓGICAS =====
        df_rallys = aplicar_regras_logicas(df_rallys)
        
        # ===== 5. VERIFICAÇÃO DE COLUNAS =====
        # Obter colunas existentes na tabela rallys
        cursor.execute("PRAGMA table_info(rallys)")
        colunas_rallys = [col[1] for col in cursor.fetchall()]
        
        # Adiciona colunas faltantes com None
        for col in colunas_rallys:
            if col not in df_rallys.columns:
                df_rallys[col] = None
        
        # Mantém apenas as colunas que existem na tabela
        df_rallys = df_rallys[[col for col in df_rallys.columns if col in colunas_rallys]]
        
        print("\n=== DADOS APÓS REGRAS LÓGICAS ===")
        print(df_rallys.head())
        
        # ===== 6. INSERIR PARTIDAS =====
        # Verificar se a tabela partidas existe
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='partidas'")
        tabela_partidas_existe = cursor.fetchone() is not None
        
        if tabela_partidas_existe:
            partidas_existentes = pd.read_sql("SELECT data, adversario FROM partidas", conn)
            df_partidas_novas = df_partidas[~df_partidas[['data', 'adversario']].apply(tuple, axis=1).isin(
                partidas_existentes[['data', 'adversario']].apply(tuple, axis=1)
            )]
        else:
            df_partidas_novas = df_partidas

        if not df_partidas_novas.empty:
            # Converter para tipos adequados
            df_partidas_novas = df_partidas_novas.astype({
                'ranking_adversario': 'Int64',
                'duracao_minutos': 'Int64',
                'cansaco_pre_jogo': 'Int64',
                'qualidade_sono': 'Int64',
                'dias_descanso': 'Int64'
            })
            
            # Obter colunas existentes na tabela partidas
            cursor.execute("PRAGMA table_info(partidas)")
            colunas_partidas = [col[1] for col in cursor.fetchall()]
            
            # Filtrar colunas do DataFrame
            df_partidas_novas = df_partidas_novas[[col for col in df_partidas_novas.columns if col in colunas_partidas]]
            
            # Usar método manual de inserção
            total_inseridos = inserir_dados_manual(conn, df_partidas_novas, "partidas", colunas_partidas)
            print(f"\n✅ {total_inseridos} nova(s) partida(s) inserida(s)")
            
            if total_inseridos > 0:
                # Backup das partidas (só após inserção bem-sucedida)
                with pd.ExcelWriter(excel_path, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                    try:
                        df_backup_partidas = pd.read_excel(excel_path, sheet_name="Backup_Partidas")
                        df_backup_partidas = pd.concat([df_backup_partidas, df_partidas_novas])
                    except:
                        df_backup_partidas = df_partidas_novas
                    df_backup_partidas.to_excel(writer, sheet_name="Backup_Partidas", index=False)
                
                # Limpa a aba Partidas
                with pd.ExcelWriter(excel_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
                    pd.DataFrame(columns=df_partidas.columns).to_excel(writer, sheet_name="Partidas", index=False)
        else:
            print("\n⏭️ Nenhuma nova partida para inserir")

        # ===== 7. INSERIR RALLYS =====
        # Verificar se a tabela rallys existe
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='rallys'")
        tabela_rallys_existe = cursor.fetchone() is not None
        
        if tabela_rallys_existe:
            rallys_existentes = pd.read_sql("SELECT partida_id, set_num, game_num, ponto_num FROM rallys", conn)
            partidas_validas = pd.read_sql("SELECT partida_id FROM partidas", conn)['partida_id'].unique()
            
            df_rallys_validos = df_rallys[df_rallys['partida_id'].isin(partidas_validas)]
            df_rallys_novos = df_rallys_validos[
                ~df_rallys_validos[['partida_id', 'set_num', 'game_num', 'ponto_num']].apply(tuple, axis=1).isin(
                    rallys_existentes[['partida_id', 'set_num', 'game_num', 'ponto_num']].apply(tuple, axis=1)
                )
            ]
        else:
            df_rallys_novos = df_rallys

        if not df_rallys_novos.empty:
            print("\n=== RALLYS A SEREM INSERIDOS ===")
            print(df_rallys_novos.head())
            
            # Usar método manual de inserção
            total_inseridos = inserir_dados_manual(conn, df_rallys_novos, "rallys", colunas_rallys)
            print(f"\n✅ {total_inseridos} novo(s) rally(s) inserido(s)")
            
            if total_inseridos > 0:
                # Backup dos rallys inseridos
                with pd.ExcelWriter(excel_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
                    try:
                        df_backup = pd.read_excel(excel_path, sheet_name="Backup_Rallys")
                        df_backup = pd.concat([df_backup, df_rallys_novos])
                    except:
                        df_backup = df_rallys_novos
                    
                    df_backup.to_excel(writer, sheet_name="Backup_Rallys", index=False)
                    pd.DataFrame(columns=colunas_rallys).to_excel(writer, sheet_name="Digitação", index=False)
            else:
                print("\n⏭️ Nenhum rally novo inserido (todos falharam na validação)")
        else:
            print("\n⏭️ Nenhum rally novo para inserir (todos já existem ou partida_id inválido)")
            # Limpa a aba mesmo sem novos dados
            with pd.ExcelWriter(excel_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
                pd.DataFrame(columns=colunas_rallys).to_excel(writer, sheet_name="Digitação", index=False)

        print("\n🔄 Processo concluído com sucesso!")

    except Exception as e:
        print(f"\n❌ ERRO GRAVE: {str(e)}")
        import traceback
        traceback.print_exc()
    finally:
        if 'conn' in locals():
            conn.close()

enviar_dados()